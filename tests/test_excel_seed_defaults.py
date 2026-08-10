import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from app import create_app
from app.extensions import db
from app.models import InventoryItem
from app.services.excel_import import seed_from_excel


class TestConfig:
    TESTING = True
    SECRET_KEY = "test"
    SQLALCHEMY_DATABASE_URI = "sqlite:///:memory:"
    SQLALCHEMY_TRACK_MODIFICATIONS = False


class ExcelSeedDefaultsTest(unittest.TestCase):
    def setUp(self):
        self.app = create_app(TestConfig)
        with self.app.app_context():
            db.create_all()

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.drop_all()

    def test_seed_applies_inventory_defaults(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "seed.xlsx"
            wb = Workbook()
            wb.remove(wb.active)

            for sheet_name, values in {
                "MASTER KONDISI": ["1 - Baik"],
                "MASTER SATUAN": ["3 - Buah"],
                "MASTER RUANG": [
                    "10531 - GEDUNG UNIT IV > GEDUNG UNIT IV LANTAI 3 > RUANG RAPAT F BIDANG PENGELOLA BMD"
                ],
                "MASTER BARANG": ["01 - Umum"],
            }.items():
                ws = wb.create_sheet(sheet_name)
                ws.append(["Nama"])
                for value in values:
                    ws.append([value])

            ws = wb.create_sheet("Worksheet")
            ws.append([
                "Nibar / Kodekib",
                "Kode Register",
                "Kode Barang",
                "Tahun Perolehan",
                "Nilai Perolehan",
                "Spesifikasi",
                "Jenis Aset",
                "Judul Buku",
                "Pencipta Buku",
                "Spesifikasi Buku",
                "Jumlah Barang",
                "Satuan Barang",
                "Status Keberadaan",
                "Jml Keberadaan (Hilang/Tdk Ditemukan)",
                "Merupakan Atribusi",
                "Nibar Atribusi",
                "Alamat",
                "Koordinat",
                "Ruangan",
                "Kondisi Barang",
                "Merk/Type",
                "Penggunaan",
                "Nama Kuasa",
                "Nama Pemakai",
                "Status Pemakai",
                "BAST",
                "Nama Dasar Penggunaan",
                "Nama Dokumen",
                "Nibar Tercatat Ganda",
                "Deskripsi Barang",
                "Keterangan",
                "Petugas (separator ';')",
                "Foto",
            ])
            ws.append([
                4001,
                "000001",
                "01 - Umum",
                2026,
                1000,
                None,
                None,
                "Judul Seed",
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                "1 - Baik",
            ])
            wb.save(path)

            with self.app.app_context():
                seed_from_excel(path)
                item = InventoryItem.query.get(4001)

                self.assertEqual(item.spesifikasi, "6-APBD")
                self.assertEqual(item.jenis_aset, "BUKU")
                self.assertEqual(item.jumlah_barang, 1)
                self.assertEqual(item.satuan_barang.nama, "3 - Buah")
                self.assertEqual(item.merupakan_atribusi, "tidak")
                self.assertEqual(item.koordinat, "-7.794439738764821, 110.36759391147048")
                self.assertEqual(item.status_pemakai, "Badan Pengelola Keuangan dan Aset DIY")
                self.assertEqual(item.nibar_tercatat_ganda, "tidak")
                self.assertIsNotNone(item.ruangan)
                self.assertEqual(item.ruangan.kode, "10531")
                self.assertEqual(item.deskripsi_barang, "Buku Umum")
                self.assertEqual(item.keterangan, "Judul Seed")


if __name__ == "__main__":
    unittest.main()
