#!/usr/bin/env python
"""Export inventory DB -> Excel v2, preserving the 38-col template and embedding photos.

Usage (from project root):
    venv/Scripts/python.exe scripts/export_v2.py
    venv/Scripts/python.exe scripts/export_v2.py --out exports/custom.xlsx

Reads the template from assets/templates/, inventory + masters from the
SQLite DB, and writes the filled workbook (with compressed foto thumbnails
in column AL) to exports/.
"""
from __future__ import annotations

import argparse
import sqlite3
import sys
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = ROOT / "assets" / "templates" / "Format_Excel_KIBE.xlsx"
DEFAULT_OUT = ROOT / "exports" / "Format_Excel_KIBE_v2.xlsx"
DEFAULT_DB = ROOT / "instance" / "inventory.db"
DEFAULT_FOTO_DIR = ROOT / "app" / "static" / "foto"

HEADER_ROW = 12
DATA_START = 13
FOTO_COLUMN = "AL"  # 0-based index 37

# Header index -> meaning (38 cols):
# 0 Nibar, 1 Kode Register, 2 Kode Barang, 3 Tahun, 4 Nilai, 5 Asal Usul,
# 6 Jenis Aset, 7 Judul, 8 Pencipta, 9 Spes Buku, 10 Asal Daerah,
# 11 Pencipta Barang, 12 Bahan, 13 Jenis Hewan, 14 Ukuran, 15 Jumlah,
# 16 Satuan, 17 Status, 18 Jml, 19 Merupakan, 20 Nibar Atr, 21 Alamat,
# 22 Koordinat, 23 Ruangan, 24 Kondisi, 25 Merk, 26 Penggunaan, 27 Kuasa,
# 28 Pemakai, 29 Status Pemakai, 30 BAST, 31 Nama Dasar, 32 Nama Dokumen,
# 33 Nibar Ganda, 34 Deskripsi, 35 Keterangan, 36 Petugas, 37 Foto


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Export inventory DB to Excel v2.")
    p.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE,
                   help=f"Seed template workbook (default: {DEFAULT_TEMPLATE})")
    p.add_argument("--out", type=Path, default=DEFAULT_OUT,
                   help=f"Output workbook path (default: {DEFAULT_OUT})")
    p.add_argument("--db", type=Path, default=DEFAULT_DB,
                   help=f"SQLite DB path (default: {DEFAULT_DB})")
    p.add_argument("--foto-dir", type=Path, default=DEFAULT_FOTO_DIR,
                   help=f"Foto directory (default: {DEFAULT_FOTO_DIR})")
    p.add_argument("--no-foto", action="store_true",
                   help="Skip embedding foto thumbnails.")
    return p.parse_args(argv)


def load_db(db_path: Path):
    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()
    masters = {}
    for tbl in ["master_barang", "master_satuan", "master_ruang", "master_kondisi"]:
        cur.execute(f"SELECT id, nama FROM {tbl}")
        masters[tbl] = {r["id"]: r["nama"] for r in cur.fetchall()}
    # For Asal Usul we don't need a master mapping, spesifikasi is already normalized.
    cur.execute("SELECT * FROM inventory_item ORDER BY nibar")
    db_rows = {r["nibar"]: dict(r) for r in cur.fetchall()}
    conn.close()
    print(f"DB rows {len(db_rows)}")
    return masters, db_rows


def update_data_rows(ws, masters, db_rows) -> dict[int, int]:
    header = [c.value for c in ws[HEADER_ROW]]
    print(f"Header 38 cols: {header[:5]} ... {header[-3:]}")

    excel_nibars = []
    for row in ws.iter_rows(min_row=DATA_START, max_row=ws.max_row, values_only=True):
        if row[0] is None:
            continue
        try:
            excel_nibars.append(int(str(row[0]).strip()))
        except (TypeError, ValueError):
            pass
    print(f"Excel data rows {len(excel_nibars)} first {excel_nibars[:3]} last {excel_nibars[-3:]}")

    db_nibars = set(db_rows.keys())
    excel_set = set(excel_nibars)
    print(f"DB nibar set {len(db_nibars)} Excel set {len(excel_set)}")
    print(f"Excel not in DB: {excel_set - db_nibars}")
    print(f"DB not in Excel: {db_nibars - excel_set}")

    nibar_to_excel_row = {}
    for idx, row in enumerate(ws.iter_rows(min_row=DATA_START, max_row=ws.max_row), start=DATA_START):
        val = row[0].value
        if val is None:
            continue
        try:
            nibar_to_excel_row[int(str(val).strip())] = idx
        except (TypeError, ValueError):
            pass

    updated = 0
    for nibar, db in db_rows.items():
        if nibar not in nibar_to_excel_row:
            new_row = ws.max_row + 1
            nibar_to_excel_row[nibar] = new_row
            print(f"Appending new NIBAR {nibar} at row {new_row}")
        r = nibar_to_excel_row[nibar]

        def set_cell(col_idx, value, number_format=None):
            cell = ws.cell(row=r, column=col_idx + 1)
            if col_idx in (0, 1):
                cell.number_format = '@'
                cell.value = str(value) if value is not None else None
            else:
                if number_format:
                    cell.number_format = number_format
                cell.value = value
            return cell

        kode_barang = masters["master_barang"].get(db["kode_barang_id"])
        satuan = masters["master_satuan"].get(db["satuan_barang_id"])
        ruangan = masters["master_ruang"].get(db["ruangan_id"])
        kondisi = masters["master_kondisi"].get(db["kondisi_barang_id"])

        set_cell(0, db["nibar"])
        set_cell(1, db["kode_register"])
        set_cell(2, kode_barang)
        set_cell(3, db["tahun_perolehan"])
        set_cell(4, db["nilai_perolehan"])
        set_cell(5, db["spesifikasi"])  # Asal Usul, already normalized (e.g. 6 - APBD)
        set_cell(6, db["jenis_aset"])
        set_cell(7, db["judul_buku"])
        set_cell(8, db["pencipta_buku"])
        set_cell(9, db["spesifikasi_buku"])
        # 10-14 (Asal Daerah .. Ukuran) have no DB mapping: keep the existing
        # Excel values untouched to preserve manual entries.
        set_cell(15, db["jumlah_barang"])
        set_cell(16, satuan)
        set_cell(17, db["status_keberadaan"])
        set_cell(18, db["jml_keberadaan"])
        set_cell(19, db["merupakan_atribusi"])
        set_cell(20, db["nibar_atribusi"])
        set_cell(21, db["alamat"])
        set_cell(22, db["koordinat"])
        set_cell(23, ruangan)
        set_cell(24, kondisi)
        set_cell(25, db["merk_type"])
        set_cell(26, db["penggunaan"])
        set_cell(27, db["nama_kuasa"])
        set_cell(28, db["nama_pemakai"])
        set_cell(29, db["status_pemakai"])
        set_cell(30, db["bast"])
        set_cell(31, db["nama_dasar_penggunaan"])
        set_cell(32, db["nama_dokumen"])
        set_cell(33, db["nibar_tercatat_ganda"])
        set_cell(34, db["deskripsi_barang"])
        set_cell(35, db["keterangan"])
        set_cell(36, db["petugas"])
        set_cell(37, None)  # Foto: image only, no directory text
        updated += 1

    print(f"Updated {updated} rows in place")
    return nibar_to_excel_row


def embed_fotos(ws, db_rows, nibar_to_excel_row, foto_dir: Path) -> int:
    ws._images = []  # clear any floating images carried over from the template
    photo_count = 0
    for nibar, db in db_rows.items():
        foto = db["foto"]
        if not foto:
            continue
        # foto is like /static/foto/1420626.jpg -> file is app/static/foto/1420626.jpg
        fpath = foto_dir / Path(foto).name
        if not fpath.exists():
            print(f"Missing foto file {fpath} for nibar {nibar}")
            continue
        r = nibar_to_excel_row.get(nibar)
        if not r:
            continue
        try:
            with PILImage.open(fpath) as im:
                if im.mode in ("RGBA", "LA", "P"):
                    im = im.convert("RGB")
                im.thumbnail((300, 300), PILImage.LANCZOS)
                bio = BytesIO()
                im.save(bio, format="JPEG", quality=80, optimize=True)
                bio.seek(0)
                xl_img = XLImage(bio)
                xl_img.width = 60
                xl_img.height = 60
                xl_img.anchor = f"{FOTO_COLUMN}{r}"
                ws.add_image(xl_img)
                ws.row_dimensions[r].height = 45
                photo_count += 1
        except Exception as e:
            print(f"Failed foto {nibar} {e}")
    print(f"Embedded {photo_count} photos")
    return photo_count


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if not args.template.exists():
        print(f"Template not found: {args.template}", file=sys.stderr)
        return 1
    if not args.db.exists():
        print(f"DB not found: {args.db}", file=sys.stderr)
        return 1

    masters, db_rows = load_db(args.db)

    wb = load_workbook(str(args.template))
    ws = wb["Worksheet"]

    nibar_to_excel_row = update_data_rows(ws, masters, db_rows)

    if not args.no_foto:
        embed_fotos(ws, db_rows, nibar_to_excel_row, args.foto_dir)
    ws.column_dimensions[FOTO_COLUMN].width = 15

    # Validations cover AA13:AA1644 etc. Only appended rows beyond the
    # template range would need validation extension (not the case today).
    print(f"Final max_row {ws.max_row} max_col {ws.max_column}")

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(args.out))
    print(f"Saved to {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
