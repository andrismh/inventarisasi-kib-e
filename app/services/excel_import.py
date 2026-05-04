from pathlib import Path
import click
from flask import current_app
from openpyxl import load_workbook

from ..extensions import db
from ..models import (
    InventoryItem,
    MasterBarang,
    MasterKondisi,
    MasterRuang,
    MasterSatuan,
)

WORKSHEET_NAME = "Worksheet"
MASTER_SHEETS = {
    "MASTER KONDISI": MasterKondisi,
    "MASTER SATUAN": MasterSatuan,
    "MASTER RUANG": MasterRuang,
    "MASTER BARANG": MasterBarang,
}


def _split_kode_nama(value):
    """Master sheets store rows like '1 - Baik' or '01.03.05.01.01.01.003 - Ilmu...'.
    The leading token before ' - ' is the kode."""
    s = str(value).strip()
    if " - " in s:
        kode, nama = s.split(" - ", 1)
        return kode.strip(), s
    return s, s


def _load_master(wb, sheet_name, model):
    ws = wb[sheet_name]
    seen = set()
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        kode, nama = _split_kode_nama(row[0])
        if kode in seen:
            continue
        seen.add(kode)
        rows.append(model(kode=kode, nama=nama))
    db.session.add_all(rows)
    db.session.flush()
    return len(rows)


def _build_lookup(model):
    return {m.nama: m.id for m in model.query.all()}


def _norm(value):
    if value is None:
        return None
    if isinstance(value, str):
        s = value.strip()
        return s if s else None
    return value


def _to_int(value):
    v = _norm(value)
    if v is None:
        return None
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def _load_worksheet(wb):
    ws = wb[WORKSHEET_NAME]
    barang = _build_lookup(MasterBarang)
    satuan = _build_lookup(MasterSatuan)
    ruang = _build_lookup(MasterRuang)
    kondisi = _build_lookup(MasterKondisi)

    inserted = 0
    skipped = 0
    seen_nibar = set()

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or row[0] is None:
            continue
        # pad to 33 columns
        cells = list(row) + [None] * (33 - len(row))

        nibar = _to_int(cells[0])
        if nibar is None or nibar in seen_nibar:
            skipped += 1
            continue
        seen_nibar.add(nibar)

        kode_barang_nama = _norm(cells[2])
        satuan_nama = _norm(cells[11])
        ruang_nama = _norm(cells[18])
        kondisi_nama = _norm(cells[19])

        kode_barang_id = barang.get(kode_barang_nama)
        satuan_barang_id = satuan.get(satuan_nama)
        ruangan_id = ruang.get(ruang_nama) if ruang_nama else None
        kondisi_barang_id = kondisi.get(kondisi_nama)

        # Required FKs missing -> skip row to keep schema invariants
        if not (kode_barang_id and satuan_barang_id and kondisi_barang_id):
            skipped += 1
            continue

        item = InventoryItem(
            nibar=nibar,
            kode_register=str(_norm(cells[1]) or ""),
            kode_barang_id=kode_barang_id,
            tahun_perolehan=_to_int(cells[3]) or 0,
            nilai_perolehan=_to_int(cells[4]) or 0,
            spesifikasi=_norm(cells[5]),
            jenis_aset=str(_norm(cells[6]) or "BUKU"),
            judul_buku=_norm(cells[7]),
            pencipta_buku=_norm(cells[8]),
            spesifikasi_buku=_norm(cells[9]),
            jumlah_barang=_to_int(cells[10]) or 1,
            satuan_barang_id=satuan_barang_id,
            status_keberadaan=_norm(cells[12]),
            jml_keberadaan=_to_int(cells[13]),
            merupakan_atribusi=_norm(cells[14]),
            nibar_atribusi=_to_int(cells[15]),
            alamat=_norm(cells[16]),
            koordinat=_norm(cells[17]) if cells[17] is None else str(_norm(cells[17])),
            ruangan_id=ruangan_id,
            kondisi_barang_id=kondisi_barang_id,
            merk_type=_norm(cells[20]),
            penggunaan=_norm(cells[21]),
            nama_kuasa=_norm(cells[22]),
            nama_pemakai=_norm(cells[23]),
            status_pemakai=_norm(cells[24]),
            bast=_norm(cells[25]),
            nama_dasar_penggunaan=_norm(cells[26]),
            nama_dokumen=_norm(cells[27]),
            nibar_tercatat_ganda=_norm(cells[28]) and str(_norm(cells[28])),
            deskripsi_barang=_norm(cells[29]),
            keterangan=_norm(cells[30]),
            petugas=_norm(cells[31]),
            foto=_norm(cells[32]) and str(_norm(cells[32])),
        )
        db.session.add(item)
        inserted += 1
        if inserted % 500 == 0:
            db.session.flush()

    db.session.flush()
    return inserted, skipped


def seed_from_excel(path: Path):
    if not path.exists():
        raise FileNotFoundError(f"Excel seed file not found at {path}")

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)

    counts = {}
    for sheet_name, model in MASTER_SHEETS.items():
        # purge existing rows so re-seeding is idempotent
        model.query.delete()
        counts[model.__name__] = _load_master(wb, sheet_name, model)

    InventoryItem.query.delete()
    inserted, skipped = _load_worksheet(wb)
    counts["InventoryItem"] = inserted
    counts["_skipped"] = skipped

    db.session.commit()
    return counts


def register_cli(app):
    @app.cli.command("seed-from-excel")
    def _cmd():
        """Seed the database from Format_Excel_KIBE.xlsx."""
        path = Path(current_app.config["EXCEL_SEED_PATH"])
        click.echo(f"Seeding from {path} ...")
        counts = seed_from_excel(path)
        for k, v in counts.items():
            label = "skipped rows" if k == "_skipped" else f"Loaded {k}"
            click.echo(f"  {label}: {v}")
        click.echo("Done.")
